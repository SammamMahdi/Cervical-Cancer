# -*- coding: utf-8 -*-
"""Builds 'Cervical Cancer Literature Review.xlsx' from lit_review_data.json.
Data is produced by a grounded+verified research workflow. This script is purely
deterministic formatting/assembly — no facts are invented here.
"""
import json, io, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"f:\Cervaical Cancer\lit_review_data.json"
OUT = r"f:\Cervaical Cancer\Cervical Cancer Literature Review.xlsx"
BUILD_DATE = "2026-06-20"

with io.open(SRC, encoding="utf-8") as f:
    data = json.load(f)

corpus     = data.get("corpus", [])
dropped    = data.get("dropped", [])
gap_panels = data.get("gapPanels", [])
field_gaps = data.get("fieldGaps", {}) or {}

# ---- styling helpers -------------------------------------------------------
HDR_FILL   = PatternFill("solid", fgColor="1F4E79")
HDR_FONT   = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=15, color="1F4E79")
SUB_FONT   = Font(bold=True, size=12, color="1F4E79")
WRAP_TOP   = Alignment(wrap_text=True, vertical="top")
CENTER     = Alignment(horizontal="center", vertical="top")
THIN       = Side(style="thin", color="BFBFBF")
BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL   = PatternFill("solid", fgColor="EFF3FA")

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER

def write_table(ws, headers, rows, widths, start_row=1, zebra=True):
    for j, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=j, value=h)
    style_header(ws, start_row, len(headers))
    r = start_row + 1
    for ridx, row in enumerate(rows):
        for j, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=j, value=val)
            cell.alignment = WRAP_TOP; cell.border = BORDER
            if zebra and ridx % 2 == 1:
                cell.fill = ALT_FILL
        r += 1
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row+len(rows)}"
    return r

wb = Workbook()

# ===== Sheet 1: README ======================================================
ws = wb.active; ws.title = "README"
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 110
def kv(ws, r, k, v, key_font=None):
    a = ws.cell(row=r, column=1, value=k); a.font = key_font or Font(bold=True)
    a.alignment = WRAP_TOP
    b = ws.cell(row=r, column=2, value=v); b.alignment = WRAP_TOP
    return r + 1
t = ws.cell(row=1, column=1, value="Cervical Cancer AI — Literature Review"); t.font = TITLE_FONT
r = 3
r = kv(ws, r, "Compiled", BUILD_DATE)
r = kv(ws, r, "Scope", "Broad cervical-cancer AI across all modalities: Pap cytology, histopathology/WSI, "
       "colposcopy, segmentation (CT/MRI/cell), tabular risk factors, multimodal fusion, plus cross-cutting "
       "techniques (knowledge distillation, federated learning, transformers, SSL/GAN augmentation).")
r = kv(ws, r, "Time range", "2021–2026, weighted toward 2024–2026 SOTA.")
r = kv(ws, r, "Papers (verified)", str(len(corpus)))
r = kv(ws, r, "How it was built",
       "Fan-out web search across 12 subfields -> per-paper field extraction from the live source -> "
       "adversarial verification: every paper's metrics & datasets were re-fetched and confirmed against the "
       "publisher/PMC/arXiv source, with a supporting quote. Numbers not stated in a source are marked "
       "'not reported'; papers that could not be confirmed real were dropped (see 'Dropped (unverified)').")
r = kv(ws, r, "Grounding note",
       "Metrics are the authors' OWN reported figures on their stated test sets; they are NOT independently "
       "reproduced and are not directly comparable across papers (different splits/datasets/protocols).")
r = kv(ws, r, "Verification legend",
       "verified = metric+datasets confirmed in source · partially-verified = paper confirmed, some field "
       "'not reported' · source-unreachable = paper confirmed via alt source.")
r += 1
s = ws.cell(row=r, column=1, value="Sheets"); s.font = SUB_FONT; r += 1
for nm, desc in [
    ("Literature Review", "Main table: methodology, novelty, training vs eval datasets, reported results, gaps, KD/FL flags, source links."),
    ("Field Gaps", "Synthesized cross-paper gaps (themes) + ranked under-explored opportunities."),
    ("Gap Panels", "Raw 3-lens gap findings: clinical translation / ML methodology / data & generalization."),
    ("Datasets Frequency", "Auto-derived count of which benchmark datasets appear across the corpus."),
    ("Dropped (unverified)", "Papers surfaced by search but NOT confirmable against a real source — excluded from the review."),
]:
    a = ws.cell(row=r, column=1, value=nm); a.font = Font(bold=True); a.alignment = WRAP_TOP
    ws.cell(row=r, column=2, value=desc).alignment = WRAP_TOP
    r += 1

# ===== Sheet 2: Literature Review ===========================================
ws = wb.create_sheet("Literature Review")
headers = ["#","Title","Authors","Year","Venue / Publisher","Modality / Subfield","Methodology",
           "Novelty / Contribution","Training Dataset(s)","Testing / Evaluation Dataset(s)",
           "Evaluation Results (reported)","Research Gaps / Limitations","KD","FL",
           "Verification","Conf.","Source Link","Code Link"]
widths  = [4,40,22,6,24,20,46,40,28,30,40,44,5,5,15,7,40,28]
rows = []
for i, p in enumerate(corpus, 1):
    rows.append([
        i, p.get("title",""), p.get("authors",""), p.get("year",""), p.get("venue",""),
        p.get("modality",""), p.get("methodology",""), p.get("novelty",""),
        p.get("training_datasets",""), p.get("eval_datasets",""), p.get("evaluation_results",""),
        p.get("gaps",""), p.get("uses_kd",""), p.get("uses_fl",""),
        p.get("_verified",""), p.get("_confidence",""), p.get("source_url",""), p.get("code_url",""),
    ])
end = write_table(ws, headers, rows, widths)
# center small cols
for col in (1,4,13,14,16):
    for rr in range(2, end):
        ws.cell(row=rr, column=col).alignment = CENTER

# ===== Sheet 3: Field Gaps ==================================================
ws = wb.create_sheet("Field Gaps")
t = ws.cell(row=1, column=1, value="Field-Level Gaps in Cervical Cancer AI (synthesized from the verified corpus)")
t.font = TITLE_FONT
themes = field_gaps.get("themes", [])
headers = ["#","Theme","Gap","Evidence (across corpus)","Research Opportunity / How to address"]
widths  = [4,26,50,46,50]
rows = [[i, th.get("theme",""), th.get("gap",""), th.get("evidence",""), th.get("opportunity","")]
        for i, th in enumerate(themes, 1)]
end = write_table(ws, headers, rows, widths, start_row=3)
# top opportunities block
r = end + 2
h = ws.cell(row=r, column=1, value="Top under-explored opportunities (ranked)"); h.font = SUB_FONT
r += 1
for k, opp in enumerate(field_gaps.get("top_opportunities", []), 1):
    c = ws.cell(row=r, column=2, value=f"{k}. {opp}"); c.alignment = WRAP_TOP
    ws.cell(row=r, column=1, value="")
    r += 1

# ===== Sheet 4: Gap Panels (raw 3-lens) =====================================
ws = wb.create_sheet("Gap Panels")
headers = ["Lens","Gap","Evidence","Why it matters"]
widths  = [34,48,46,48]
rows = []
for panel in gap_panels:
    lens = panel.get("lens","")
    for g in panel.get("gaps", []):
        rows.append([lens, g.get("gap",""), g.get("evidence",""), g.get("why_it_matters","")])
write_table(ws, headers, rows, widths)

# ===== Sheet 5: Datasets Frequency (auto-derived) ===========================
ws = wb.create_sheet("Datasets Frequency")
known = ["SipakMed","Herlev","LBC","Liquid-Based Cytology","CRIC","Cervix93","UCI","Mendeley",
         "ISBI","IARC","Intel & MobileODT","Kaggle","HErlev","BCNB","TCGA","private","in-house",
         "colposcopy","whole slide","WSI","CT","MRI","PET"]
counts = {}
blob = []
for p in corpus:
    blob.append((p.get("training_datasets","") + " || " + p.get("eval_datasets","")).lower())
for name in known:
    n = sum(1 for b in blob if name.lower() in b)
    if n:
        counts[name] = n
rows = sorted(counts.items(), key=lambda kv: -kv[1])
write_table(ws, ["Dataset / source (keyword)","# papers in corpus"], [[k, v] for k, v in rows], [40, 22])
ws.cell(row=1, column=4, value="(Auto-derived by keyword scan of dataset fields; indicative, not exhaustive.)").font = Font(italic=True, color="808080")

# ===== Sheet 6: Dropped =====================================================
ws = wb.create_sheet("Dropped (unverified)")
ws.cell(row=1, column=1, value="Papers surfaced by search but NOT confirmable against a real source — EXCLUDED.").font = SUB_FONT
write_table(ws, ["Title","Source URL","Verifier note"], [[d.get("title",""), d.get("source_url",""), d.get("note","")] for d in dropped], [50,40,60], start_row=3)

wb.save(OUT)
print(f"WROTE {OUT}")
print(f"  Literature Review rows : {len(corpus)}")
print(f"  Field Gap themes       : {len(field_gaps.get('themes', []))}")
print(f"  Gap-panel findings     : {sum(len(p.get('gaps', [])) for p in gap_panels)}")
print(f"  Dropped (unverified)   : {len(dropped)}")
print(f"  Sheets                 : {wb.sheetnames}")
